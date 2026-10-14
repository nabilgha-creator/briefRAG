"""aperçu HTML des documents pour le lecteur (docx, xlsx, html) et chemins des sources, sans base"""
import pytest
from docx import Document as DocumentWord
from openpyxl import Workbook

from backend.chargement.apercu import DocumentIntrouvable, apercu_html, chemin_du_document


def test_docx_titres_paragraphes_et_listes(tmp_path) :
    document = DocumentWord()
    document.add_heading("Étape 1 — Qualifier le défaut", level=1)
    paragraphe = document.add_paragraph("Demander au client ")
    paragraphe.add_run("le code affiché").bold = True
    document.add_paragraph("Vérifier le filtre", style="List Number")
    document.add_paragraph("Relancer une vidange", style="List Number")
    document.add_paragraph("Conclusion")
    chemin = tmp_path / "procedure.docx"
    document.save(chemin)

    html = apercu_html(chemin)
    assert "<h2>Étape 1 — Qualifier le défaut</h2>" in html
    assert "<p>Demander au client <strong>le code affiché</strong></p>" in html
    assert "<ol>\n<li>Vérifier le filtre</li>\n<li>Relancer une vidange</li>\n</ol>\n<p>Conclusion</p>" in html


def test_xlsx_titre_en_paragraphe_puis_tableau(tmp_path) :
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Gamme WX"
    feuille.append(["Référentiel des codes erreur — Lave-linge série WX"])
    feuille.append(["Code", "Signification"])
    feuille.append(["E23", "Défaut pompe de vidange"])
    chemin = tmp_path / "referentiel.xlsx"
    classeur.save(chemin)

    html = apercu_html(chemin)
    assert html.startswith("<h2>Gamme WX</h2>\n<p>Référentiel des codes erreur")
    assert "<th>Code</th><th>Signification</th>" in html and "<td>E23</td><td>Défaut pompe de vidange</td>" in html


def test_html_sans_script_ni_evenement(tmp_path) :
    chemin = tmp_path / "faq.html"
    chemin.write_text("<html><body><script>alert(1)</script><main><h1 onclick='x()'>FAQ</h1>"
                      "<details><summary>Question ?</summary><p>Réponse.</p></details></main></body></html>", encoding="utf-8")
    html = apercu_html(chemin)
    assert "<h1>FAQ</h1>" in html and "<summary>Question ?</summary>" in html
    assert "script" not in html and "onclick" not in html


def test_pas_d_apercu_pour_un_pdf(tmp_path) :
    chemin = tmp_path / "manuel.pdf"
    chemin.write_bytes(b"%PDF-1.4")
    with pytest.raises(ValueError) :
        apercu_html(chemin)


@pytest.mark.parametrize("source", ["../.env", "../../backend/config.py", "Manuels/inexistant.pdf"])
def test_source_hors_du_dossier_ou_absente_refusee(source) :
    with pytest.raises(DocumentIntrouvable) :
        chemin_du_document(source)
